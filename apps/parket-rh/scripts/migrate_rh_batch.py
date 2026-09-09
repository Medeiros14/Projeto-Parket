#!/usr/bin/env python3
"""
Migração BATCH do export Convenia → schema rh.

Em vez de 3000 queries individuais, monta poucos INSERTs grandes com
multi-row VALUES. Roda em segundos.

Uso:
  export SUPABASE_MGMT_TOKEN=sbp_xxx
  python3 migrate_rh_batch.py [--xlsx PATH]
"""
import argparse
import os
import re
import sys
import unicodedata
from datetime import datetime, date
from typing import Any, Optional

import requests
from openpyxl import load_workbook

SUPABASE_REF = "hbxpilrxmitvzebluoom"
MGMT_API = f"https://api.supabase.com/v1/projects/{SUPABASE_REF}/database/query"
TOKEN = os.environ.get("SUPABASE_MGMT_TOKEN") or os.environ.get("SB_TOKEN")


def sql(query: str) -> Any:
    if not TOKEN:
        sys.exit("ERRO: defina SUPABASE_MGMT_TOKEN")
    r = requests.post(
        MGMT_API,
        headers={"Authorization": f"Bearer {TOKEN}", "Content-Type": "application/json"},
        json={"query": query},
        timeout=180,
    )
    if r.status_code >= 400:
        raise RuntimeError(f"SQL falhou ({r.status_code}): {r.text[:600]}\n--- query (head) ---\n{query[:600]}")
    return r.json()


def esc(v: Any) -> str:
    if v is None or v == "":
        return "NULL"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, (date, datetime)):
        return f"'{v.isoformat()[:10]}'"
    s = str(v).replace("'", "''")
    return f"'{s}'"


def n(s) -> Optional[str]:
    if s is None: return None
    s = str(s).strip()
    if not s or s.lower() in ("não informado", "nao informado", "n/a", "-", "—", "null"):
        return None
    return re.sub(r"\s+", " ", s)


def parse_date(s) -> Optional[str]:
    if s is None or s == "": return None
    if isinstance(s, datetime): return s.date().isoformat()
    if isinstance(s, date): return s.isoformat()
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try: return datetime.strptime(s, fmt).date().isoformat()
        except ValueError: continue
    return None


def parse_int(s) -> Optional[int]:
    if s is None or s == "": return None
    try: return int(float(str(s).replace(",", ".").strip()))
    except: return None


def parse_float(s) -> Optional[float]:
    if s is None or s == "": return None
    try: return float(str(s).replace(",", ".").strip())
    except: return None


def clean_cpf(s) -> Optional[str]:
    if s is None: return None
    digits = re.sub(r"\D", "", str(s))
    return digits if len(digits) == 11 else None


def clean_cnpj(s) -> Optional[str]:
    if s is None: return None
    return re.sub(r"\D", "", str(s)) or None


def parse_bool(v) -> Optional[bool]:
    if v is None: return None
    s = str(v).strip().lower()
    if s in ("sim", "s", "yes", "true", "1"): return True
    if s in ("não", "nao", "n", "no", "false", "0"): return False
    return None


def slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.upper().strip()


def values(rows: list[tuple]) -> str:
    """rows = [(esc'd_str, esc'd_str, ...), ...]"""
    return ",\n  ".join("(" + ",".join(r) + ")" for r in rows)


# ─── Main ──────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="/tmp/rh_completo.xlsx")
    args = ap.parse_args()

    print(f"📂 Abrindo {args.xlsx}…")
    wb = load_workbook(args.xlsx, data_only=True)

    # ── Phase 1: Empresas (PARKET + EXCLUSIVE) ──
    print("\n🏢 Resolvendo empresas…")
    empresas = {
        "22951220000133": {"nome": "PARKET", "cnpj_fmt": "22.951.220/0001-33"},
        "51992934000198": {"nome": "EXCLUSIVE", "cnpj_fmt": "51.992.934/0001-98"},
    }
    for cnpj, info in empresas.items():
        rows = sql(f"SELECT id FROM core.empresas WHERE regexp_replace(cnpj, '\\D', '', 'g') = '{cnpj}' LIMIT 1")
        if rows:
            info["id"] = rows[0]["id"]
        else:
            r = sql(
                f"INSERT INTO core.empresas (cnpj, razao_social, nome_fantasia, ativo) "
                f"VALUES ({esc(info['cnpj_fmt'])}, {esc(info['nome'])}, {esc(info['nome'])}, true) RETURNING id"
            )
            info["id"] = r[0]["id"]
        print(f"  {info['nome']} ({cnpj}) → {info['id']}")

    def emp_id_for(cnpj_raw: str) -> Optional[str]:
        c = clean_cnpj(cnpj_raw)
        return empresas.get(c, {}).get("id") if c else None

    # ── Phase 2: Coletar catálogos do worksheet ──
    print("\n📚 Coletando catálogos…")
    ws = wb["Worksheet"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    g = lambda row, key: row[idx[key]] if key in idx else None

    all_rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"  {len(all_rows)} colaboradores no worksheet")

    depts_set = set()       # (empresa_id, nome)
    times_set = set()       # (empresa_id, nome, dept_nome)
    cargos_set = {}         # (empresa_id, nome) -> {salario, cbo, descricao}
    ccs_set = set()         # (empresa_id, nome)
    sind_set = {}           # cnpj_or_slug -> {nome, cnpj_fmt, website, telefone, data_base}

    for row in all_rows:
        nome_col = n(g(row, "Nome do colaborador"))
        if not nome_col: continue
        eid = emp_id_for(g(row, "Dados da empresa - CNPJ"))
        if not eid: continue
        d = n(g(row, "Informações de trabalho - Departamento"))
        t = n(g(row, "Informações de trabalho - Time"))
        c = n(g(row, "Informações de trabalho - Cargo"))
        cc = n(g(row, "Informações de trabalho - Centro de custo"))
        if d: depts_set.add((eid, d))
        if t: times_set.add((eid, t, d))
        if c:
            cargos_set.setdefault((eid, c), {
                "salario": parse_float(g(row, "Informações de trabalho - Salário")),
                "cbo": n(g(row, "Informações de trabalho - CBO")),
                "descricao": n(g(row, "Informações de trabalho - Descrição do cargo")),
            })
        if cc: ccs_set.add((eid, cc))

        sn = n(g(row, "Informações de trabalho - Sindicato"))
        if sn:
            cnpj_s = clean_cnpj(g(row, "Informações de trabalho - Sindicato - CNPJ"))
            key = cnpj_s or slug(sn)
            if key not in sind_set:
                sind_set[key] = {
                    "nome": sn,
                    "cnpj_fmt": (f"{cnpj_s[:2]}.{cnpj_s[2:5]}.{cnpj_s[5:8]}/{cnpj_s[8:12]}-{cnpj_s[12:14]}"
                                 if cnpj_s else None),
                    "website": n(g(row, "Informações de trabalho - Sindicato - Website")),
                    "telefone": n(g(row, "Informações de trabalho - Sindicato - Telefone")),
                    "data_base": n(g(row, "Informações de trabalho - Sindicato - Database")),
                }

    print(f"  {len(depts_set)} departamentos, {len(times_set)} times, {len(cargos_set)} cargos, "
          f"{len(ccs_set)} centros de custo, {len(sind_set)} sindicatos")

    # ── Phase 3: Bulk INSERT catálogos ──
    print("\n💾 Inserindo catálogos…")

    # Departamentos
    if depts_set:
        vs = [(esc(eid), esc(nome)) for (eid, nome) in depts_set]
        sql(f"INSERT INTO rh.departamentos (empresa_id, nome) VALUES\n  {values(vs)}\n"
            f"ON CONFLICT (empresa_id, nome) DO UPDATE SET nome = EXCLUDED.nome")
    deps_rows = sql("SELECT id, empresa_id, nome FROM rh.departamentos")
    dept_map = {(r["empresa_id"], slug(r["nome"])): r["id"] for r in deps_rows}

    # Times (refs deptos)
    if times_set:
        vs = []
        for (eid, t_nome, d_nome) in times_set:
            d_id = dept_map.get((eid, slug(d_nome))) if d_nome else None
            vs.append((esc(eid), esc(d_id), esc(t_nome)))
        sql(f"INSERT INTO rh.times (empresa_id, departamento_id, nome) VALUES\n  {values(vs)}\n"
            f"ON CONFLICT (empresa_id, nome) DO UPDATE SET nome = EXCLUDED.nome")
    times_rows = sql("SELECT id, empresa_id, nome FROM rh.times")
    time_map = {(r["empresa_id"], slug(r["nome"])): r["id"] for r in times_rows}

    # CBOs existentes (FK)
    cbo_rows = sql("SELECT codigo FROM rh.cbos")
    cbo_codes = {r["codigo"] for r in cbo_rows}

    # Cargos
    if cargos_set:
        vs = []
        for (eid, c_nome), info in cargos_set.items():
            cbo = info["cbo"] if info["cbo"] in cbo_codes else None
            vs.append((
                esc(eid), esc(c_nome), esc(info["descricao"]),
                esc(info["salario"]), esc(cbo)
            ))
        sql(f"INSERT INTO rh.cargos (empresa_id, nome, descricao, salario_base_default, cbo_codigo) VALUES\n  {values(vs)}\n"
            f"ON CONFLICT (empresa_id, nome) DO UPDATE SET nome = EXCLUDED.nome")
    cargos_rows = sql("SELECT id, empresa_id, nome FROM rh.cargos")
    cargo_map = {(r["empresa_id"], slug(r["nome"])): r["id"] for r in cargos_rows}

    # Centros de custo
    if ccs_set:
        vs = [(esc(eid), esc(cc)) for (eid, cc) in ccs_set]
        sql(f"INSERT INTO rh.centros_custo_rh (empresa_id, nome) VALUES\n  {values(vs)}\n"
            f"ON CONFLICT (empresa_id, nome) DO UPDATE SET nome = EXCLUDED.nome")
    ccs_rows = sql("SELECT id, empresa_id, nome FROM rh.centros_custo_rh")
    cc_map = {(r["empresa_id"], slug(r["nome"])): r["id"] for r in ccs_rows}

    # Sindicatos (sem unique em (cnpj) NOT NULL → SELECT depois)
    if sind_set:
        vs = []
        for info in sind_set.values():
            vs.append((
                esc(info["nome"]), esc(info["cnpj_fmt"]),
                esc(info["website"]), esc(info["telefone"]), esc(info["data_base"])
            ))
        # Tem UNIQUE(cnpj) — quando cnpj null, dup ok
        sql(f"INSERT INTO rh.sindicatos (nome, cnpj, website, telefone, data_base) VALUES\n  {values(vs)}\n"
            f"ON CONFLICT (cnpj) DO UPDATE SET nome = EXCLUDED.nome")
    sind_rows = sql("SELECT id, nome, cnpj FROM rh.sindicatos")
    sind_map = {}
    for r in sind_rows:
        cnpj_clean = clean_cnpj(r["cnpj"])
        if cnpj_clean: sind_map[cnpj_clean] = r["id"]
        sind_map[slug(r["nome"])] = r["id"]

    print(f"  ✓ {len(dept_map)} dept, {len(time_map)} times, {len(cargo_map)} cargos, "
          f"{len(cc_map)} ccs, {len(sind_rows)} sindicatos")

    # ── Phase 4: Bulk INSERT colaboradores ──
    print("\n👥 Inserindo colaboradores…")

    colab_data = []  # cada: (cols dict + meta {empresa_id, contrato_data, dept, time, cargo, cc, sind})
    for row in all_rows:
        nome_c = n(g(row, "Nome do colaborador"))
        if not nome_c: continue
        eid = emp_id_for(g(row, "Dados da empresa - CNPJ"))
        if not eid: continue

        cpf = clean_cpf(g(row, "Documentos - CPF"))
        ext_id = n(g(row, "Dados da empresa - ID (identificador do colaborador)"))

        col = {
            "external_id_convenia": ext_id,
            "nome": nome_c,
            "nome_social": n(g(row, "Informações pessoais - Nome social")),
            "cpf": cpf,
            "data_nascimento": parse_date(g(row, "Informações pessoais - Data de nascimento")),
            "sexo": n(g(row, "Informações pessoais - Gênero")),
            "genero_documento": n(g(row, "Informações pessoais - Gênero no documento")),
            "estado_civil": n(g(row, "Informações pessoais - Estado civil")),
            "raca_cor": n(g(row, "Informações pessoais - Cor/Raça")),
            "uf_natal": n(g(row, "Informações pessoais - UF natal")),
            "cidade_natal": n(g(row, "Informações pessoais - Cidade natal")),
            "nome_mae": n(g(row, "Informações pessoais - Nome da mãe")),
            "nome_pai": n(g(row, "Informações pessoais - Nome do pai")),
            "email_pessoal": n(g(row, "Contatos - E-mail pessoal")),
            "email_profissional": n(g(row, "Informações de trabalho - E-mail profissional")),
            "celular": n(g(row, "Contatos - Celular")),
            "telefone_residencial": n(g(row, "Contatos - Telefone")),
            "endereco_logradouro": n(g(row, "Endereços - Endereço")),
            "endereco_numero": n(g(row, "Endereços - Número")),
            "endereco_complemento": n(g(row, "Endereços - Complemento do endereço")),
            "endereco_bairro": n(g(row, "Endereços - Bairro")),
            "endereco_cidade": n(g(row, "Endereços - Cidade")),
            "endereco_uf": n(g(row, "Endereços - Estado")),
            "endereco_cep": n(g(row, "Endereços - CEP")),
            "rg_numero": n(g(row, "Documentos - RG - Número")),
            "rg_data_emissao": parse_date(g(row, "Documentos - RG - Data de emissão")),
            "rg_orgao_emissor": n(g(row, "Documentos - RG - Órgão emissor")),
            "rg_uf_emissor": n(g(row, "Documentos - RG - UF emissor")),
            "ctps_numero": n(g(row, "Documentos - CTPS - Número")),
            "ctps_serie": n(g(row, "Documentos - CTPS - Número de série")),
            "ctps_data_emissao": parse_date(g(row, "Documentos - CTPS - Data de emissão")),
            "ctps_uf_emissora": n(g(row, "Documentos - CTPS - Estado emissor")),
            "pis_pasep": n(g(row, "Documentos - CTPS - PIS")),
            "titulo_eleitor_numero": n(g(row, "Documentos - Titulo de eleitor - Número")),
            "titulo_eleitor_zona": n(g(row, "Documentos - Titulo de eleitor - Ala Eleitoral")),
            "titulo_eleitor_secao": n(g(row, "Documentos - Titulo de eleitor - Sessão")),
            "titulo_eleitor_uf": n(g(row, "Documentos - Titulo de eleitor - Estado")),
            "titulo_eleitor_cidade": n(g(row, "Documentos - Titulo de eleitor - Cidade")),
            "reservista_numero": n(g(row, "Documentos - Reservista - Número")),
            "reservista_ra": n(g(row, "Documentos - Reservista - RA")),
            "reservista_serie": n(g(row, "Documentos - Reservista - Série")),
            "cnh_numero": n(g(row, "Documentos - CNH - Número")),
            "cnh_data_emissao": parse_date(g(row, "Documentos - CNH - Data de emissão")),
            "cnh_validade": parse_date(g(row, "Documentos - CNH - Validade")),
            "cnh_categoria": n(g(row, "Documentos - CNH - Categoria")),
            "deficiencia_tipo": n(g(row, "Deficiência - Tipo")),
            "deficiencia_observacoes": n(g(row, "Deficiência - Observações")),
            "qtd_dependentes_ir": parse_int(g(row, "Quantidade de dependentes declarados no IR")) or 0,
            "status_dados": "completo" if cpf else "incompleto",
        }

        meta = {
            "empresa_id": eid,
            "matricula": (str(parse_int(g(row, "Informações de trabalho - Matrícula")))
                          if parse_int(g(row, "Informações de trabalho - Matrícula")) is not None
                          else n(g(row, "Informações de trabalho - Matrícula"))),
            "data_admissao": parse_date(g(row, "Informações de trabalho - Data de admissão")),
            "data_termino_contrato": parse_date(g(row, "Informações de trabalho - Data de término do contrato de aprendiz")),
            "vinculo": n(g(row, "Informações de trabalho - Vínculo")),
            "tipo_admissao": n(g(row, "Informações de trabalho - Tipo de admissão")),
            "salario_base": parse_float(g(row, "Informações de trabalho - Salário")) or 0,
            "tipo_salario": n(g(row, "Informações de trabalho - Tipo de salário")),
            "forma_pagamento": n(g(row, "Informações de trabalho - Forma de pagamento")),
            "hora_contratual": parse_float(g(row, "Informações de trabalho - Hora contratual")),
            "horas_mensais": parse_float(g(row, "Horas mensais")),
            "regime_jornada": n(g(row, "Regime de jornada do colaborador")),
            "jornada_descricao": n(g(row, "Jornada de trabalho - Jornada de trabalho")),
            "jornada_motivo": n(g(row, "Jornada de trabalho - Motivo da jornada")),
            "jornada_observacoes": n(g(row, "Jornada de trabalho - Observações")),
            "periodo_experiencia_dias": parse_int(g(row, "Período de experiência - Dias totais")),
            "primeiro_termino_experiencia": parse_date(g(row, "Período de experiência - Primeiro término")),
            "segundo_termino_experiencia": parse_date(g(row, "Período de experiência - Segundo término")),
            "data_exame_admissional": parse_date(g(row, "Informações de trabalho - Data do exame admissional")),
            "cargo_confianca": parse_bool(g(row, "Informações de trabalho - Cargo de confiança")) or False,
            "primeiro_emprego": parse_bool(g(row, "Informações de trabalho - Primeiro emprego")) or False,
            "aposentado": parse_bool(g(row, "Informações de trabalho - Aposentado")) or False,
            "seguro_desemprego": parse_bool(g(row, "Informações de trabalho - Tem seguro desemprego?")),
            "cipa": parse_bool(g(row, "Informações de trabalho - CIPA")) or False,
            "registro_ponto": parse_bool(g(row, "Informações de trabalho - Possui registro de ponto?")),
            "estabilidade": n(g(row, "Informações de trabalho - Estabilidade")),
            "inscricao_orgao_classe": n(g(row, "Informações de trabalho - Inscrição no órgão de classe")),
            "conselho_profissional": n(g(row, "Informações de trabalho - Conselho profissional")),
            "senioridade": n(g(row, "Informações de trabalho - Senioridade")),
            "nivel_senioridade": n(g(row, "Informações de trabalho - Nível de senioridade")),
            "data_demissao": parse_date(g(row, "Desligados - Data de desligamento")),
            "motivo_afastamento": n(g(row, "Desligados - Motivo do desligamento")),
            "dept_slug": slug(n(g(row, "Informações de trabalho - Departamento")) or ""),
            "time_slug": slug(n(g(row, "Informações de trabalho - Time")) or ""),
            "cargo_slug": slug(n(g(row, "Informações de trabalho - Cargo")) or ""),
            "cc_slug": slug(n(g(row, "Informações de trabalho - Centro de custo")) or ""),
            "sind_key": (clean_cnpj(g(row, "Informações de trabalho - Sindicato - CNPJ"))
                         or (slug(n(g(row, "Informações de trabalho - Sindicato")) or "")
                             if n(g(row, "Informações de trabalho - Sindicato")) else None)),
            "gestor_nome": n(g(row, "Informações de trabalho - Gestor")),
        }
        colab_data.append((col, meta))

    # Insert colaboradores em 1 SQL com RETURNING
    cols_order = [
        "external_id_convenia", "nome", "nome_social", "cpf", "data_nascimento",
        "sexo", "genero_documento", "estado_civil", "raca_cor",
        "uf_natal", "cidade_natal", "nome_mae", "nome_pai",
        "email_pessoal", "email_profissional", "celular", "telefone_residencial",
        "endereco_logradouro", "endereco_numero", "endereco_complemento",
        "endereco_bairro", "endereco_cidade", "endereco_uf", "endereco_cep",
        "rg_numero", "rg_data_emissao", "rg_orgao_emissor", "rg_uf_emissor",
        "ctps_numero", "ctps_serie", "ctps_data_emissao", "ctps_uf_emissora", "pis_pasep",
        "titulo_eleitor_numero", "titulo_eleitor_zona", "titulo_eleitor_secao",
        "titulo_eleitor_uf", "titulo_eleitor_cidade",
        "reservista_numero", "reservista_ra", "reservista_serie",
        "cnh_numero", "cnh_data_emissao", "cnh_validade", "cnh_categoria",
        "deficiencia_tipo", "deficiencia_observacoes", "qtd_dependentes_ir",
        "status_dados",
    ]
    vs = []
    for col, _meta in colab_data:
        vs.append(tuple(esc(col[k]) for k in cols_order))
    print(f"  inserindo {len(vs)} colaboradores em 1 query…")
    sql_text = (
        f"INSERT INTO rh.colaboradores ({', '.join(cols_order)}, observacoes) VALUES\n  "
        + ",\n  ".join("(" + ",".join(r) + f", 'Importado de Convenia em {datetime.utcnow().isoformat()}Z')" for r in vs)
        + "\nON CONFLICT (external_id_convenia) DO UPDATE SET "
        + ", ".join(f"{k} = EXCLUDED.{k}" for k in cols_order if k != "external_id_convenia")
        + "\nRETURNING id, external_id_convenia, nome"
    )
    res = sql(sql_text)
    # Map external_id → colab_id e nome→colab_id (case-insensitive) pra side tables
    colab_by_ext = {r["external_id_convenia"]: r["id"] for r in res if r["external_id_convenia"]}
    colab_by_nome = {slug(r["nome"]): r["id"] for r in res}
    print(f"  ✓ {len(res)} colaboradores")

    # Resolver CPF→colab_id pra side tables (tem CPFs únicos no schema)
    cpf_rows = sql("SELECT id, cpf, nome FROM rh.colaboradores WHERE cpf IS NOT NULL")
    colab_by_cpf = {r["cpf"]: r["id"] for r in cpf_rows}
    # Atualiza colab_by_nome com TODOS os colabs (caso o RETURNING não tenha retornado todos por conflict)
    all_rows_db = sql("SELECT id, nome, external_id_convenia FROM rh.colaboradores")
    for r in all_rows_db:
        colab_by_nome[slug(r["nome"])] = r["id"]
        if r["external_id_convenia"]:
            colab_by_ext[r["external_id_convenia"]] = r["id"]

    # ── Phase 5: Contratos ──
    print("\n📝 Inserindo contratos…")
    ctr_vs = []
    for col, m in colab_data:
        cid = colab_by_ext.get(col["external_id_convenia"]) \
              or colab_by_cpf.get(col["cpf"]) \
              or colab_by_nome.get(slug(col["nome"]))
        if not cid:
            print(f"  ⚠️ colab sem id: {col['nome']}")
            continue
        ctr_vs.append((
            esc(cid), esc(m["empresa_id"]),
            esc(m["matricula"]),
            esc(cargo_map.get((m["empresa_id"], m["cargo_slug"])) if m["cargo_slug"] else None),
            esc(dept_map.get((m["empresa_id"], m["dept_slug"])) if m["dept_slug"] else None),
            esc(time_map.get((m["empresa_id"], m["time_slug"])) if m["time_slug"] else None),
            esc(cc_map.get((m["empresa_id"], m["cc_slug"])) if m["cc_slug"] else None),
            esc(sind_map.get(m["sind_key"]) if m["sind_key"] else None),
            esc(m["salario_base"]),
            esc(m["tipo_salario"]), esc(m["forma_pagamento"]), esc(m["hora_contratual"]),
            esc(m["data_admissao"]), esc(m["data_termino_contrato"]),
            esc(m["data_exame_admissional"]),
            esc(m["periodo_experiencia_dias"]),
            esc(m["primeiro_termino_experiencia"]), esc(m["segundo_termino_experiencia"]),
            esc(m["vinculo"]), esc(m["tipo_admissao"]),
            esc(m["jornada_descricao"]), esc(m["jornada_motivo"]), esc(m["jornada_observacoes"]),
            esc(m["regime_jornada"]), esc(m["horas_mensais"]),
            esc(m["cargo_confianca"]), esc(m["estabilidade"]),
            esc(m["primeiro_emprego"]), esc(m["aposentado"]),
            esc(m["seguro_desemprego"]), esc(m["cipa"]),
            esc(m["registro_ponto"] if m["registro_ponto"] is not None else True),
            esc(m["inscricao_orgao_classe"]), esc(m["conselho_profissional"]),
            esc(m["senioridade"]), esc(m["nivel_senioridade"]),
            esc(m["data_demissao"]), esc(m["motivo_afastamento"]),
            esc("desligado" if m["data_demissao"] else "ativo"),
            esc(m.get("gestor_nome")),
        ))

    ctr_cols = [
        "colaborador_id", "empresa_id", "matricula",
        "cargo_id", "departamento_id", "time_id", "centro_custo_rh_id", "sindicato_id",
        "salario_base",
        "tipo_salario", "forma_pagamento", "hora_contratual",
        "data_admissao", "data_termino_contrato",
        "data_exame_admissional",
        "periodo_experiencia_dias",
        "primeiro_termino_experiencia", "segundo_termino_experiencia",
        "vinculo", "tipo_admissao",
        "jornada_descricao", "jornada_motivo", "jornada_observacoes",
        "regime_jornada", "horas_mensais",
        "cargo_confianca", "estabilidade",
        "primeiro_emprego", "aposentado",
        "seguro_desemprego", "cipa",
        "registro_ponto",
        "inscricao_orgao_classe", "conselho_profissional",
        "senioridade", "nivel_senioridade",
        "data_demissao", "motivo_afastamento",
        "status",
        "gestor_nome_snapshot",
    ]
    print(f"  inserindo {len(ctr_vs)} contratos em 1 query…")
    sql(f"INSERT INTO rh.contratos ({', '.join(ctr_cols)}) VALUES\n  {values(ctr_vs)}\n"
        f"ON CONFLICT (empresa_id, matricula) DO NOTHING")
    nctr = sql("SELECT count(*)::int AS c FROM rh.contratos")[0]["c"]
    print(f"  ✓ {nctr} contratos no total")

    # ── Phase 6: Side tables ──
    def colab_id_lookup(empresa_nome, nome, cpf):
        cid = colab_by_cpf.get(clean_cpf(cpf)) if cpf else None
        if not cid: cid = colab_by_nome.get(slug(nome or ""))
        return cid

    # Nacionalidades
    print("\n🌍 Nacionalidades…")
    ws_n = wb["Nacionalidades"]
    h_n = [c.value for c in ws_n[1]]
    inx = {h: i for i, h in enumerate(h_n)}
    nat_vs = []
    for r in ws_n.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf = r[inx["CPF do colaborador"]]
        nat = n(r[inx["Informações pessoais - Nacionalidade"]])
        if not (nome and nat): continue
        cid = colab_id_lookup(None, nome, cpf)
        if not cid: continue
        nat_vs.append((esc(cid), esc(nat), "1"))
    if nat_vs:
        sql(f"INSERT INTO rh.colaborador_nacionalidades (colaborador_id, nacionalidade, ordem) VALUES\n  {values(nat_vs)}\n"
            f"ON CONFLICT (colaborador_id, nacionalidade) DO NOTHING")
    print(f"  ✓ {len(nat_vs)} nacionalidades")

    # Contatos emergência
    print("\n📞 Contatos de emergência…")
    ws_e = wb["Contatos de emergência"]
    h_e = [c.value for c in ws_e[1]]
    inx = {h: i for i, h in enumerate(h_e)}
    em_vs = []
    for r in ws_e.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf = r[inx["CPF do colaborador"]]
        c_nome = n(r[inx["Contatos de emergência - Nome"]])
        if not (nome and c_nome): continue
        cid = colab_id_lookup(None, nome, cpf)
        if not cid: continue
        em_vs.append((
            esc(cid), esc(c_nome),
            esc(n(r[inx["Contatos de emergência - Relação"]])),
            esc(n(r[inx["Contatos de emergência - Telefone"]])),
            esc(n(r[inx["Contatos de emergência - Telefone comercial"]])),
            esc(n(r[inx["Contatos de emergência - Celular"]])),
            esc(n(r[inx["Contatos de emergência - E-mail"]])),
            "1",
        ))
    if em_vs:
        sql(f"INSERT INTO rh.contatos_emergencia "
            f"(colaborador_id, nome, relacao, telefone, telefone_comercial, celular, email, ordem) VALUES\n  {values(em_vs)}")
    print(f"  ✓ {len(em_vs)} contatos")

    # Dados bancários
    print("\n💳 Dados bancários…")
    ws_b = wb["Dados bancários"]
    h_b = [c.value for c in ws_b[1]]
    inx = {h: i for i, h in enumerate(h_b)}
    bk_vs = []
    for r in ws_b.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf = r[inx["CPF do colaborador"]]
        if not nome: continue
        cid = colab_id_lookup(None, nome, cpf)
        if not cid: continue
        tipo_raw = (n(r[inx["Dados bancários - Tipo de conta"]]) or "").lower()
        tipo = "corrente" if "corrente" in tipo_raw else ("poupanca" if "poupan" in tipo_raw else
               ("salario" if "salár" in tipo_raw else None))
        bk_vs.append((
            esc(cid),
            esc(n(r[inx["Dados bancários - Banco"]])),
            esc(tipo),
            esc(n(r[inx["Dados bancários - Modalidade"]])),
            esc(n(r[inx["Dados bancários - Agência bancária"]])),
            esc(n(r[inx["Dados bancários - Conta bancária"]])),
            esc(n(r[inx["Dados bancários - Dígito"]])),
            esc(n(r[inx["Dados bancários - PIX"]])),
            "true",  # principal
        ))
    if bk_vs:
        sql(f"INSERT INTO rh.dados_bancarios "
            f"(colaborador_id, banco_nome, tipo, modalidade, agencia, conta, digito, pix_chave, principal) VALUES\n  {values(bk_vs)}")
    print(f"  ✓ {len(bk_vs)} contas bancárias")

    # Formação
    print("\n🎓 Formação acadêmica…")
    ws_f = wb["Formação acadêmica"]
    h_f = [c.value for c in ws_f[1]]
    inx = {h: i for i, h in enumerate(h_f)}
    fo_vs = []
    for r in ws_f.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf = r[inx["CPF do colaborador"]]
        if not nome: continue
        cid = colab_id_lookup(None, nome, cpf)
        if not cid: continue
        ano_raw = r[inx["Formação acadêmica - Ano de conclusão"]]
        try: ano = int(str(ano_raw).strip()[:4]) if ano_raw else None
        except: ano = None
        fo_vs.append((
            esc(cid),
            esc(n(r[inx["Formação acadêmica - Escolaridade"]])),
            esc(n(r[inx["Formação acadêmica - Instituição de ensino"]])),
            esc(n(r[inx["Formação acadêmica - Curso"]])),
            esc(ano),
        ))
    if fo_vs:
        sql(f"INSERT INTO rh.formacao_academica "
            f"(colaborador_id, escolaridade, instituicao, curso, ano_conclusao) VALUES\n  {values(fo_vs)}")
    print(f"  ✓ {len(fo_vs)} formações")

    # Dependentes
    print("\n👨‍👩‍👧 Dependentes…")
    ws_d = wb["Informações do dependente"]
    h_d = [c.value for c in ws_d[1]]
    inx = {h: i for i, h in enumerate(h_d)}
    de_vs = []
    for r in ws_d.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf_t = r[inx["CPF do colaborador"]]
        d_nome = n(r[inx["Informações de dependentes - Nome"]])
        if not (nome and d_nome): continue
        cid = colab_id_lookup(None, nome, cpf_t)
        if not cid: continue
        de_vs.append((
            esc(cid), esc(d_nome),
            esc(clean_cpf(r[inx["Informações de dependentes - CPF"]])),
            esc(parse_date(r[inx["Informações de dependentes - Data de aniversário"]])),
            esc(n(r[inx["Informações de dependentes - Relação"]])),
            esc(parse_bool(r[inx["Informações de dependentes - Incluir para fins de imposto de renda"]]) or False),
            esc(parse_bool(r[inx["Informações de dependentes - Renda familiar"]])),
            esc(n(r[inx["Informações de dependentes - Nome da mãe"]])),
            esc(n(r[inx["Informações de dependentes - Telefone"]])),
            esc(n(r[inx["Informações de dependentes - E-mail"]])),
            esc(n(r[inx["Informações de dependentes - Descrição"]])),
            esc(parse_bool(r[inx["Informações de dependentes - Dependente possui incapacidade física ou mental para trabalho?"]])),
            esc(parse_bool(r[inx["Informações de dependentes - Estrangeiro"]])),
            esc(n(r[inx["Informações de dependentes - Gênero no documento"]])),
            esc(n(r[inx["Informações de dependentes - Escolaridade"]])),
        ))
    if de_vs:
        sql(f"INSERT INTO rh.dependentes "
            f"(colaborador_id, nome, cpf, data_nascimento, parentesco, irrf, renda_familiar, "
            f" nome_mae, telefone, email, descricao, incapacidade_fisica_mental, estrangeiro, "
            f" genero_documento, escolaridade) VALUES\n  {values(de_vs)}")
    print(f"  ✓ {len(de_vs)} dependentes")

    # Histórico cargos/salários
    print("\n📈 Histórico cargos/salários…")
    ws_h = wb["Histórico cargos e salários"]
    h_h = [c.value for c in ws_h[1]]
    inx = {h: i for i, h in enumerate(h_h)}

    # precisa achar contrato_id por colab+empresa
    ctr_rows = sql("SELECT id, colaborador_id, empresa_id FROM rh.contratos")
    ctr_by_colab = {}
    for cr in ctr_rows:
        ctr_by_colab.setdefault(cr["colaborador_id"], cr["id"])  # 1º contrato

    hi_vs = []
    for r in ws_h.iter_rows(min_row=2, values_only=True):
        nome = n(r[inx["Nome do colaborador"]])
        cpf_t = r[inx["CPF do colaborador"]]
        d_de = parse_date(r[inx["De"]])
        if not (nome and d_de): continue
        cid = colab_id_lookup(None, nome, cpf_t)
        if not cid: continue
        ctr_id = ctr_by_colab.get(cid)
        if not ctr_id: continue
        hi_vs.append((
            esc(ctr_id), esc(d_de),
            esc(parse_date(r[inx["Até"]])),
            esc(n(r[inx["Vínculo"]])),
            esc(n(r[inx["Categoria de trabalhadores"]])),
            esc(n(r[inx["Cargo"]])),
            esc(n(r[inx["Departamento"]])),
            esc(n(r[inx["Time"]])),
            esc(n(r[inx["Centro de custo"]])),
            esc(parse_float(r[inx["Salário"]])),
            esc(n(r[inx["Senioridade"]])),
            esc(n(r[inx["Nível de senioridade"]])),
            esc(n(r[inx["Motivo"]])),
            esc(n(r[inx["Descrição"]])),
        ))
    if hi_vs:
        sql(f"INSERT INTO rh.contrato_alteracoes "
            f"(contrato_id, data_de, data_ate, vinculo, categoria_trabalhador, "
            f" cargo_snapshot, departamento_snapshot, time_snapshot, centro_custo_snapshot, "
            f" salario, senioridade, nivel_senioridade, motivo, descricao) VALUES\n  {values(hi_vs)}")
    print(f"  ✓ {len(hi_vs)} alterações")

    # ── Resumo ──
    print("\n" + "=" * 60)
    n_colab = sql("SELECT count(*)::int AS c FROM rh.colaboradores")[0]["c"]
    n_ctr = sql("SELECT count(*)::int AS c FROM rh.contratos")[0]["c"]
    n_ctr_emp = sql("""SELECT e.razao_social, count(*)::int AS c FROM rh.contratos c
                       JOIN core.empresas e ON e.id = c.empresa_id GROUP BY e.razao_social""")
    print(f"✅ {n_colab} colaboradores, {n_ctr} contratos")
    for row in n_ctr_emp:
        print(f"   {row['razao_social']}: {row['c']}")


if __name__ == "__main__":
    main()
